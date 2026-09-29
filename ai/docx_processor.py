from docx import Document


def extract_docx_text(file_path):
    """
    Extract readable text from a DOCX file.

    Extracts:
    - Paragraphs
    - Tables

    Returns:
        str: Extracted document text
    """

    document = Document(file_path)
    text_parts = []

    # Extract normal paragraphs
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    # Extract tables
    for table in document.tables:
        for row in table.rows:
            row_data = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                if cell_text:
                    row_data.append(cell_text)

            if row_data:
                text_parts.append(" | ".join(row_data))

    return "\n".join(text_parts)


if __name__ == "__main__":
    file_path = input("Enter DOCX file path: ").strip()

    extracted_text = extract_docx_text(file_path)

    print("\nEXTRACTED DOCX TEXT")
    print("===================")
    print(extracted_text)